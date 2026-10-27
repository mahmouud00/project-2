"""Materials from an Excel workbook: preview, import (all or nothing), prices converted to the calculation unit,
supplier discounts, and a company never changing MYTech's library rows."""
import io
from decimal import Decimal as D

from openpyxl import Workbook

HEAD = ["السيستم", "الكود", "الاسم", "النوع", "الوحدة", "طول العود (مم)", "المجموعة", "اللون", "كود الصنف (SKU)",
        "سعر المورد", "السعر على", "ملاحظات", "وزن المتر (كجم)", "خصم المورد %"]


def _book(rows, systems=(("XSYS", "سيستم تجريبي", "upvc", ""),), colors=(("WH", "أبيض", "#FFFFFF"), ("BG", "بيج", ""))):
    wb = Workbook()
    ws = wb.active
    ws.title = "الألوان"
    ws.append(["الكود", "اللون", "اللون (hex)"])
    for c in colors:
        ws.append(list(c))
    s = wb.create_sheet("السيستمات")
    s.append(["الكود", "الاسم", "النوع", "ملاحظات"])
    for x in systems:
        s.append(list(x))
    m = wb.create_sheet("قطاعات")
    m.append(HEAD)
    for r in rows:
        m.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, h, data, apply):
    return client.post(f"/api/v1/library/import-excel?apply={str(apply).lower()}", headers=h,
                       files={"file": ("m.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_import_preview_then_apply(tenant_a, client):
    h = tenant_a.h
    rows = [
        ["XSYS", "X-H01", "حلق مفصلي", "قطاع", "متر", 6000, "مفصلي", "أبيض", "X-H01-WH", 177.07, "المتر", "", "", 10],
        ["XSYS", "X-H01", "حلق مفصلي", "قطاع", "متر", 6000, "مفصلي", "بيج", "X-H01-BG", 177.07, "المتر", "", "", 10],
        ["", "X-SL", "شريحة شاتر", "قطاع", "عود", 6000, "شرائح", "أبيض", "", 48.40, "العود (6 م)", "", "", ""],
        ["", "X-GD", "مجرى شريحة", "قطاع", "كجم", 6000, "مجاري", "", "", 380, "الكيلو", "", 0.3333, ""],
        ["", "X-ACC", "مقبض", "إكسسوار", "قطعة", "", "مقابض", "", "", 50.72, "القطعة", "بدون خصم", "", 0],
    ]
    data = _book(rows)
    pre = _post(client, h, data, False)
    assert pre.status_code == 200, pre.text
    p = pre.json()
    assert p["materials"]["new"] == 4 and p["systems"]["new"] == 1 and p["colors"]["new"] == 2 and p["errors"] == []
    assert client.get("/api/v1/library/materials?q=X-H01", headers=h).json()["items"] == []   # preview saves nothing
    r = _post(client, h, data, True)
    assert r.status_code == 200, r.text
    mats = {m["code"]: m for m in client.get("/api/v1/library/materials?q=X-", headers=h).json()["items"]}
    assert set(mats) == {"X-H01", "X-SL", "X-GD", "X-ACC"}
    assert {v["sku"] for v in mats["X-H01"]["variants"]} == {"X-H01-WH", "X-H01-BG"}
    assert mats["X-SL"]["unit"] == "m" and mats["X-GD"]["unit"] == "m"          # bar and kg prices -> per metre
    from app.db.session import tenant_session
    from app.modules.sales.pricing import CostResolver
    with tenant_session(tenant_a.id) as db:
        res = CostResolver(db)
        assert res.resolve("X-H01", "WH")[0] == D("159.3630")          # 177.07 − 10 %
        assert res.resolve("X-SL", "WH")[0] == D("8.0667")             # 48.40 per 6 m bar
        assert res.resolve("X-GD", None)[0] == D("126.6540")           # 380 per kg × 0.3333 kg/m
        assert res.resolve("X-ACC", None)[0] == D("50.72")
    # importing again updates, never duplicates
    again = _post(client, h, data, True).json()
    assert again["materials"]["updated"] == 4 and again["materials"]["new"] == 0


def test_import_is_all_or_nothing(tenant_a, client):
    data = _book([["XSYS", "Y-1", "صنف", "قطاع", "متر", 6000, "", "أبيض", "", 10, "", "", "", ""],
                  ["XSYS", "Y-2", "صنف", "حاجة غريبة", "متر", 6000, "", "أحمر", "", 10, "", "", "", ""]])
    r = _post(client, tenant_a.h, data, True)
    assert r.status_code == 422 and r.json()["error"]["code"] == "IMPORT_ERRORS"
    errors = r.json()["error"]["details"]["errors"]
    assert "سطر 3" in errors[0]["where"]
    assert client.get("/api/v1/library/materials?q=Y-1", headers=tenant_a.h).json()["items"] == []


def test_company_prices_library_material_without_changing_it(platform, tenant_a, tenant_b, client):
    from tests.library_helpers import build_library
    lib = build_library(client, platform, code="XLB")
    data = _book([["", "XLB-FRAME", "اسم الشركة", "قطاع", "متر", 6000, "", "أبيض", "", 99, "", "", "", ""]],
                 systems=(), colors=(("W", "أبيض", ""),))
    r = _post(client, tenant_a.h, data, True)
    assert r.status_code == 200, r.text
    assert r.json()["materials"]["library"] == 1 and r.json()["materials"]["new"] == 0
    m = client.get("/api/v1/library/materials?q=XLB-FRAME", headers=tenant_a.h).json()["items"][0]
    assert m["name"] != "اسم الشركة"                                    # the library row is MYTech's
    assert [(p["color_code"], p["cost"]) for p in m["company_prices"]] == [("W", "99.0000")]   # shown to the company
    other = client.get("/api/v1/library/materials?q=XLB-FRAME", headers=tenant_b.h).json()["items"][0]
    assert other["company_prices"] == []                                                       # never to another one
    from app.db.session import tenant_session
    from app.modules.sales.pricing import CostResolver
    with tenant_session(tenant_a.id) as db:
        assert CostResolver(db).resolve("XLB-FRAME", "W")[0] == D("99")
    with tenant_session(tenant_b.id) as db:
        assert CostResolver(db).resolve("XLB-FRAME", "W")[0] != D("99")  # another company keeps its own price
    assert lib


def test_template_download(tenant_a, client):
    r = client.get("/api/v1/library/import-excel/template", headers=tenant_a.h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
