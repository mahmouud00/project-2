"""استيراد الخامات من إكسيل: profile systems, colours, materials with their colours and the supplier prices.

The workbook has a sheet "الألوان" (الكود، اللون، hex), a sheet "السيستمات" (الكود، الاسم، النوع، ملاحظات) and one
or more material sheets whose first row has these headers (any order; extra columns are ignored):

    السيستم · الكود · الاسم · النوع · الوحدة · طول العود (مم) · المجموعة · اللون · كود الصنف (SKU) · سعر المورد ·
    السعر على · ملاحظات · وزن المتر (كجم) · خصم المورد %

One row = one material in one colour. Rows are matched by material code and colour: new ones are added, existing ones
updated (never deleted). Who imports decides where it goes:
  * MYTech (platform)  -> the shared library; the price becomes the reference price.
  * a company          -> its own materials when the code is not in the library; the price (and discount) becomes the
                          company's price for that material and colour. Library rows are never changed by a company.

Prices are converted to the unit the calculations use: a price per bar (عود) becomes a price per metre
(÷ bar length in metres); a price per kg of a profile with a weight per metre becomes a price per metre (× kg/m).
These are unit conversions of money, not deductions.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.audit import audit
from app.core.errors import ValidationFailed
from app.db.base import PLATFORM_COMPANY_ID
from app.modules.library.models import Color, CompanyMaterialSetting, Material, MaterialVariant, ProfileSystem
from app.modules.library.models import SystemConstant

KIND = {"قطاع": "profile", "حديد تدعيم": "reinforcement", "حديد": "reinforcement", "إكسسوار": "accessory",
        "اكسسوار": "accessory", "كاوتش": "gasket", "مستهلكات": "consumable", "بانل": "panel", "زجاج": "glass",
        "أخرى": "other"}
UNIT = {"متر": "m", "م": "m", "قطعة": "pcs", "طقم": "set", "لفة": "roll", "م²": "m2", "م2": "m2", "كجم": "kg",
        "كيلو": "kg", "لتر": "l", "عود": "bar"}
SYSTEM_TYPE = {"upvc": "upvc", "shutter": "shutter", "aluminum": "aluminum", "شاتر": "shutter", "ألومنيوم": "aluminum"}
HEAD = {"السيستم": "system", "الكود": "code", "الاسم": "name", "النوع": "kind", "الوحدة": "unit",
        "طول العود (مم)": "bar_length_mm", "المجموعة": "category", "اللون": "color", "كود الصنف (SKU)": "sku",
        "سعر المورد": "price", "السعر على": "basis", "ملاحظات": "notes", "وزن المتر (كجم)": "weight",
        "خصم المورد %": "discount"}


@dataclass
class Report:
    systems: dict = field(default_factory=lambda: {"new": 0, "updated": 0})
    colors: dict = field(default_factory=lambda: {"new": 0})
    materials: dict = field(default_factory=lambda: {"new": 0, "updated": 0, "library": 0})
    variants: dict = field(default_factory=lambda: {"new": 0, "updated": 0})
    prices: int = 0
    rows: int = 0
    errors: list = field(default_factory=list)

    def as_dict(self) -> dict:
        return {"systems": self.systems, "colors": self.colors, "materials": self.materials, "variants": self.variants,
                "prices": self.prices, "rows": self.rows, "errors": self.errors[:200], "error_count": len(self.errors)}


def _dec(v) -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"رقم مش صحيح: {v}")


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def read_workbook(data: bytes) -> tuple[list[dict], list[dict], list[tuple[str, int, dict]]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValidationFailed("The file is not an Excel workbook (.xlsx)", code="BAD_EXCEL")
    colors, systems, rows = [], [], []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            continue
        cols = [_text(h) for h in header]
        if ws.title.strip() == "الألوان":
            for r in it:
                if r and _text(r[0]):
                    colors.append({"code": _text(r[0]), "name": _text(r[1]) if len(r) > 1 else "",
                                   "hex": _text(r[2]) if len(r) > 2 else ""})
        elif ws.title.strip() == "السيستمات":
            for r in it:
                if r and _text(r[0]):
                    systems.append({"code": _text(r[0]), "name": _text(r[1]), "type": _text(r[2]) if len(r) > 2 else "",
                                    "notes": _text(r[3]) if len(r) > 3 else ""})
        elif "الكود" in cols and "الاسم" in cols:
            index = {HEAD[c]: i for i, c in enumerate(cols) if c in HEAD}
            for n, r in enumerate(it, start=2):
                if not r or not any(v is not None and _text(v) for v in r):
                    continue
                rows.append((ws.title, n, {k: (r[i] if i < len(r) else None) for k, i in index.items()}))
    return colors, systems, rows


def import_workbook(db, ctx, data: bytes, apply: bool) -> dict:
    """Preview (apply=False: nothing is saved) or import. Everything or nothing."""
    colors, systems, rows = read_workbook(data)
    rep = Report()
    platform = ctx.company_id == PLATFORM_COMPANY_ID
    mine = ctx.company_id

    color_by_name: dict[str, str] = {}
    known_colors = {c.code: c for c in db.scalars(select(Color))}
    for c in colors:
        color_by_name[c["name"]] = c["code"]
        if c["code"] not in known_colors:
            rep.colors["new"] += 1
            if apply:
                known_colors[c["code"]] = Color(code=c["code"], name=c["name"] or c["code"], hex=c["hex"] or None)
                db.add(known_colors[c["code"]])
    for code, c in known_colors.items():
        color_by_name.setdefault(c.name, code)

    system_ids: dict[str, object] = {}
    for s in systems:
        existing = db.scalar(select(ProfileSystem).where(ProfileSystem.code == s["code"],
                                                         ProfileSystem.company_id.in_([mine, PLATFORM_COMPANY_ID])))
        if existing is not None:
            system_ids[s["code"]] = existing.id
            if existing.company_id == mine:
                rep.systems["updated"] += 1
                if apply:
                    existing.name, existing.description = s["name"] or existing.name, s["notes"] or existing.description
            continue
        rep.systems["new"] += 1
        system_ids[s["code"]] = None        # known (made on import)
        if apply:
            mtype = SYSTEM_TYPE.get(s["type"].lower(), "upvc")
            row = ProfileSystem(code=s["code"], name=s["name"] or s["code"], material_type=mtype, unit="mm",
                                description=s["notes"], status="trial", supplier="")
            db.add(row)
            db.flush()
            for name, label in (("install_clearance_width", "خلوص التركيب في العرض"),
                                ("install_clearance_height", "خلوص التركيب في الارتفاع")):
                db.add(SystemConstant(system_id=row.id, name=name, label=label, value=Decimal(0), sort=-1))
            audit(db, row, "create")
            system_ids[s["code"]] = row.id

    seen_materials: dict[str, Material | None] = {}
    for sheet, n, r in rows:
        rep.rows += 1
        where = f"{sheet} سطر {n}"
        try:
            code, name = _text(r.get("code")), _text(r.get("name"))
            if not code or not name:
                raise ValueError("الكود والاسم مطلوبين")
            kind = KIND.get(_text(r.get("kind")), None)
            if kind is None:
                raise ValueError(f"النوع مش معروف: {_text(r.get('kind'))}")
            unit = UNIT.get(_text(r.get("unit")), None)
            if unit is None:
                raise ValueError(f"الوحدة مش معروفة: {_text(r.get('unit'))}")
            bar = _dec(r.get("bar_length_mm"))
            weight = _dec(r.get("weight"))
            price = _dec(r.get("price"))
            discount = _dec(r.get("discount"))
            if discount is not None and not (0 <= discount <= 100):
                raise ValueError("الخصم لازم بين 0 و 100")
            color_name = _text(r.get("color"))
            color = color_by_name.get(color_name) if color_name else ""
            if color is None:
                raise ValueError(f"اللون مش موجود في شيت الألوان: {color_name}")
            sku = _text(r.get("sku")) or (f"{code}-{color}" if color else code)
            system_code = _text(r.get("system"))
            system_id = system_ids.get(system_code)
            if system_code and system_code not in system_ids:
                found = db.scalar(select(ProfileSystem).where(ProfileSystem.code == system_code))
                if found is None:
                    raise ValueError(f"السيستم مش موجود: {system_code}")
                system_id = system_ids[system_code] = found.id
            # money in the unit the calculations use
            if price is not None and unit == "bar":
                if not bar:
                    raise ValueError("سعر العود محتاج طول العود")
                price = (price / (bar / 1000)).quantize(Decimal("0.0001"))
            if unit == "bar":
                unit = "m"
            if price is not None and unit == "kg" and weight and kind in ("profile", "reinforcement"):
                price, unit = (price * weight).quantize(Decimal("0.0001")), "m"
        except ValueError as e:
            rep.errors.append({"where": where, "message": str(e)})
            continue

        if code not in seen_materials:
            own = db.scalar(select(Material).where(Material.code == code, Material.company_id == mine))
            library = None if platform or own else db.scalar(
                select(Material).where(Material.code == code, Material.company_id == PLATFORM_COMPANY_ID))
            fields = dict(name=name, kind=kind, unit=unit, bar_length_mm=bar, weight_kg_per_m=weight,
                          category=_text(r.get("category")), system_id=system_id)
            if own is not None:
                rep.materials["updated"] += 1
                if apply:
                    for k, v in fields.items():
                        setattr(own, k, v)
                    own.version += 1
                material = own
            elif library is not None:
                rep.materials["library"] += 1       # a company prices a library material: the row stays MYTech's
                material = library
            else:
                rep.materials["new"] += 1
                material = None
                if apply:
                    material = Material(code=code, status="approved", **fields)
                    db.add(material)
                    db.flush()
                    audit(db, material, "create")
            seen_materials[code] = material
        material = seen_materials[code]
        if not apply:
            if color:
                rep.variants["new"] += 1
            if price is not None or discount is not None:
                rep.prices += 1
            continue

        variant = None
        if color:
            variant = db.scalar(select(MaterialVariant).where(MaterialVariant.material_id == material.id,
                                                              MaterialVariant.color_code == color))
            if variant is None and material.company_id == mine:
                variant = MaterialVariant(material_id=material.id, color_code=color, sku=sku, name=name)
                db.add(variant)
                rep.variants["new"] += 1
            elif variant is not None and variant.company_id == mine:
                variant.sku, variant.is_active = sku, True
                rep.variants["updated"] += 1
        if price is None and discount is None:
            continue
        rep.prices += 1
        if platform:
            if variant is not None:
                variant.reference_cost = price
            else:
                material.reference_cost = price
            continue
        setting = db.scalar(select(CompanyMaterialSetting).where(CompanyMaterialSetting.material_id == material.id,
                                                                 CompanyMaterialSetting.color_code == color))
        if setting is None:
            setting = CompanyMaterialSetting(material_id=material.id, color_code=color)
            db.add(setting)
        if price is not None:
            setting.unit_cost = price
        if discount is not None:
            setting.discount_pct = discount
    if apply and rep.errors:
        db.rollback()
        raise ValidationFailed("Some rows have errors: nothing was imported", code="IMPORT_ERRORS",
                               details=rep.as_dict())
    return rep.as_dict()
